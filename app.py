import os
import io
import json
import uuid
import hashlib
import warnings
import traceback
from datetime import datetime

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from flask import Flask, request, jsonify, send_file, session
from flask_cors import CORS
import re

warnings.filterwarnings("ignore")

app = Flask(__name__, static_folder=".", template_folder=".")
app.secret_key = os.environ.get("SECRET_KEY", "audit-secret-key-2024")
CORS(app)

# In-memory store (works on Render single worker)
_REPORT_STORE = {}

# ── colour palette ──────────────────────────────────────────────────────────
C = {
    "bg_dark":    "1A0A0A",
    "bg_header":  "2D0808",
    "accent":     "C0392B",
    "accent2":    "E74C3C",
    "gold":       "D4AC0D",
    "text_light": "F5F5F5",
    "text_dim":   "AAAAAA",
    "row_alt":    "2A1010",
    "green":      "27AE60",
    "orange":     "E67E22",
    "red":        "C0392B",
    "white":      "FFFFFF",
    "border":     "5C1A1A",
}

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, size=11, color="F5F5F5", italic=False):
    return Font(bold=bold, size=size, color=color, italic=italic, name="Calibri")

def _border():
    side = Side(border_style="thin", color=C["border"])
    return Border(left=side, right=side, top=side, bottom=side)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


# ── helpers ─────────────────────────────────────────────────────────────────
def detect_column_type(series):
    s = series.dropna()
    if len(s) == 0:
        return "empty"
    if pd.api.types.is_numeric_dtype(s):
        return "numeric"
    try:
        pd.to_datetime(s, infer_datetime_format=True, errors="raise")
        return "date"
    except Exception:
        pass
    return "text"


def find_salary_cols(df):
    kw = ["salary", "sal", "wage", "pay", "ctc", "compensation", "income"]
    return [c for c in df.columns
            if any(k in c.lower() for k in kw) and pd.api.types.is_numeric_dtype(df[c])]


def find_profit_cols(df):
    kw = ["profit", "margin", "earnings", "net", "gain", "revenue", "sales"]
    return [c for c in df.columns
            if any(k in c.lower() for k in kw) and pd.api.types.is_numeric_dtype(df[c])]


def check_business_logic(df):
    issues = []
    rev_cols  = [c for c in df.columns if "revenue" in c.lower() and pd.api.types.is_numeric_dtype(df[c])]
    cost_cols = [c for c in df.columns if "cost"    in c.lower() and pd.api.types.is_numeric_dtype(df[c])]
    profit_cols = find_profit_cols(df)

    if rev_cols and cost_cols and profit_cols:
        r, co, p = rev_cols[0], cost_cols[0], profit_cols[0]
        expected = df[r] - df[co]
        diff = (df[p] - expected).abs()
        bad = diff[diff > 0.01]
        if not bad.empty:
            issues.append({
                "type": "logic_error", "severity": "high",
                "description": f"Profit mismatch: {p} ≠ {r} − {co}",
                "affected_rows": bad.index.tolist()[:20], "count": len(bad)
            })

    disc_cols = [c for c in df.columns if "discount" in c.lower() and pd.api.types.is_numeric_dtype(df[c])]
    for dc in disc_cols:
        bad = df[df[dc] > 100]
        if not bad.empty:
            issues.append({
                "type": "invalid_value", "severity": "high",
                "description": f"Discount > 100% in column '{dc}'",
                "affected_rows": bad.index.tolist()[:20], "count": len(bad)
            })

    age_cols = [c for c in df.columns if c.lower() == "age" and pd.api.types.is_numeric_dtype(df[c])]
    for ac in age_cols:
        bad = df[(df[ac] < 0) | (df[ac] > 120)]
        if not bad.empty:
            issues.append({
                "type": "invalid_value", "severity": "medium",
                "description": f"Unrealistic age values in '{ac}'",
                "affected_rows": bad.index.tolist()[:20], "count": len(bad)
            })
    return issues


def detect_statistical_anomalies(df):
    anomalies = []
    for col in df.select_dtypes(include=[np.number]).columns:
        s = df[col].dropna()
        if len(s) < 10:
            continue
        mean, std = s.mean(), s.std()
        if std == 0:
            continue
        z = ((s - mean) / std).abs()
        outliers = z[z > 3]
        if not outliers.empty:
            anomalies.append({
                "column": col, "count": len(outliers),
                "rows": outliers.index.tolist()[:10],
                "values": s[outliers.index].tolist()[:10]
            })
    return anomalies


def detect_suspicious_edits(df):
    flags = []
    for col in df.select_dtypes(include=[np.number]).columns:
        s = df[col].dropna()
        if len(s) < 5:
            continue
        round_mask = s.apply(lambda x: x % 1000 == 0 and x != 0)
        round_rows = s[round_mask]
        ratio = len(round_rows) / len(s)
        if 0 < ratio < 0.15 and len(round_rows) >= 2:
            flags.append({
                "column": col, "type": "suspiciously_round",
                "description": f"{len(round_rows)} suspiciously round values (multiples of 1000) in '{col}'",
                "rows": round_rows.index.tolist()[:10]
            })
    return flags


def full_audit(df, filename):
    report = {
        "filename": filename,
        "timestamp": datetime.now().isoformat(),
        "shape": {"rows": len(df), "cols": len(df.columns)},
        "columns": list(df.columns),
        "col_types": {c: detect_column_type(df[c]) for c in df.columns},
        "missing": {}, "duplicates": {}, "negative_profits": {},
        "salary_issues": {}, "business_logic_errors": [],
        "statistical_anomalies": [], "suspicious_edits": [], "kpis": {},
    }

    missing = df.isnull().sum()
    report["missing"] = {
        col: {"count": int(cnt), "pct": round(cnt / len(df) * 100, 2)}
        for col, cnt in missing.items() if cnt > 0
    }

    dup_mask = df.duplicated()
    report["duplicates"] = {
        "count": int(dup_mask.sum()),
        "rows": df[dup_mask].index.tolist()[:20]
    }

    for col in find_profit_cols(df):
        neg = df[df[col] < 0]
        if not neg.empty:
            report["negative_profits"][col] = {
                "count": len(neg),
                "rows": neg.index.tolist()[:20],
                "min": float(neg[col].min())
            }

    for col in find_salary_cols(df):
        s = df[col].dropna()
        report["salary_issues"][col] = {
            "total": float(s.sum()), "mean": float(s.mean()),
            "median": float(s.median()), "min": float(s.min()),
            "max": float(s.max()), "zero_count": int((s == 0).sum()),
            "negative_count": int((s < 0).sum())
        }

    report["business_logic_errors"]  = check_business_logic(df)
    report["statistical_anomalies"]  = detect_statistical_anomalies(df)
    report["suspicious_edits"]       = detect_suspicious_edits(df)

    num_df = df.select_dtypes(include=[np.number])
    kpis = {}
    for col in num_df.columns:
        s = num_df[col].dropna()
        if len(s) == 0:
            continue
        kpis[col] = {
            "sum": round(float(s.sum()), 2), "mean": round(float(s.mean()), 2),
            "median": round(float(s.median()), 2), "std": round(float(s.std()), 2),
            "min": round(float(s.min()), 2), "max": round(float(s.max()), 2),
            "count": int(s.count()), "null_count": int(num_df[col].isnull().sum())
        }
    report["kpis"] = kpis

    total_issues = (
        len(report["missing"]) +
        report["duplicates"]["count"] +
        sum(v["count"] for v in report["negative_profits"].values()) +
        len(report["business_logic_errors"]) +
        len(report["statistical_anomalies"]) +
        len(report["suspicious_edits"])
    )
    report["health_score"] = (
        100 if total_issues == 0 else
        85  if total_issues < 5  else
        65  if total_issues < 15 else
        40  if total_issues < 30 else 20
    )
    return report


# ── Excel builders ───────────────────────────────────────────────────────────
def write_section_title(ws, row, col, text, col_span=6):
    cell = ws.cell(row=row, column=col, value=text)
    cell.fill = _fill(C["bg_header"])
    cell.font = _font(bold=True, size=13, color=C["accent2"])
    cell.alignment = _left()
    ws.merge_cells(start_row=row, start_column=col,
                   end_row=row, end_column=col + col_span - 1)


def build_cover_sheet(wb, audit, df):
    ws = wb.create_sheet("Audit Summary")
    ws.sheet_properties.tabColor = "C0392B"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 50
    ws.merge_cells("A2:D2")
    t = ws["A2"]
    t.value = "EXCEL AUDIT & REPORTING SYSTEM"
    t.fill = _fill(C["accent"])
    t.font = Font(bold=True, size=22, color=C["white"], name="Calibri")
    t.alignment = _center()

    ws.row_dimensions[4].height = 28
    ws.merge_cells("A4:D4")
    s = ws["A4"]
    s.value = f"File: {audit['filename']}   |   Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}"
    s.fill = _fill(C["bg_dark"])
    s.font = _font(italic=True, size=11, color=C["text_dim"])
    s.alignment = _center()

    ws.row_dimensions[6].height = 60
    metrics = [
        ("Total Rows",    audit["shape"]["rows"],  C["accent"]),
        ("Total Columns", audit["shape"]["cols"],  C["bg_header"]),
        ("Health Score",  f"{audit['health_score']}%",
         C["green"] if audit["health_score"] > 70 else C["orange"] if audit["health_score"] > 40 else C["red"]),
        ("Issues Found",
         sum([len(audit["missing"]), audit["duplicates"]["count"],
              len(audit["business_logic_errors"]), len(audit["statistical_anomalies"])]),
         C["orange"]),
    ]
    for i, (label, val, color) in enumerate(metrics, 1):
        c = ws.cell(row=6, column=i)
        c.value = f"{label}\n{val}"
        c.fill = _fill(color)
        c.font = Font(bold=True, size=14, color=C["white"], name="Calibri")
        c.alignment = _center()
        c.border = _border()

    row = 8
    write_section_title(ws, row, 1, "  ISSUE BREAKDOWN", 4)
    row += 1
    for ci, h in enumerate(["Category", "Details", "Severity", "Status"], 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = _fill(C["bg_header"])
        c.font = _font(bold=True, color=C["accent2"])
        c.alignment = _center()
        c.border = _border()
        ws.row_dimensions[row].height = 20
    row += 1

    issues_data = []
    if audit["missing"]:
        issues_data.append(("Missing Values",
            f"{sum(v['count'] for v in audit['missing'].values())} cells across {len(audit['missing'])} columns",
            "MEDIUM", "Review"))
    if audit["duplicates"]["count"] > 0:
        issues_data.append(("Duplicate Rows",
            f"{audit['duplicates']['count']} duplicate rows detected", "HIGH", "Fix"))
    for col, info in audit["negative_profits"].items():
        issues_data.append(("Negative Profits",
            f"{info['count']} rows in '{col}' (min: {info['min']:,.2f})", "HIGH", "Fix"))
    for err in audit["business_logic_errors"]:
        issues_data.append(("Business Logic", err["description"], err["severity"].upper(), "Fix"))
    for a in audit["statistical_anomalies"]:
        issues_data.append(("Statistical Anomaly",
            f"{a['count']} outliers in '{a['column']}'", "MEDIUM", "Review"))
    for s in audit["suspicious_edits"]:
        issues_data.append(("Suspicious Edit", s["description"], "LOW", "Note"))
    if not issues_data:
        issues_data.append(("No Issues", "Dataset passed all checks", "NONE", "Clean"))

    sev_color = {"HIGH": C["red"], "MEDIUM": C["orange"], "LOW": C["gold"], "NONE": C["green"]}
    for alt, (cat, det, sev, status) in enumerate(issues_data):
        bg = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
        ws.row_dimensions[row].height = 22
        for ci, val in enumerate([cat, det, sev, status], 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.fill = _fill(sev_color.get(sev, C["bg_dark"])) if ci == 3 else _fill(bg)
            c.font = _font(bold=True, size=10, color=C["white"]) if ci == 3 else _font(size=10)
            c.alignment = _left() if ci == 2 else _center()
            c.border = _border()
        row += 1
    return ws


def build_cleaned_data_sheet(wb, df, audit):
    ws = wb.create_sheet("Cleaned Data")
    ws.sheet_properties.tabColor = "27AE60"
    ws.freeze_panes = "A2"

    dup_rows = set(audit["duplicates"]["rows"])
    for ci, col in enumerate(df.columns, 1):
        try:
            col_str_len = df[col].fillna("").astype(str).str.len().max()
            col_str_len = int(col_str_len) if pd.notna(col_str_len) else 10
        except Exception:
            col_str_len = 10
        ws.column_dimensions[get_column_letter(ci)].width = min(max(len(str(col)), col_str_len) + 4, 30)

    ws.row_dimensions[1].height = 28
    for ci, col in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=ci, value=col)
        c.fill = _fill(C["accent"])
        c.font = _font(bold=True, size=11, color=C["white"])
        c.alignment = _center()
        c.border = _border()

    null_fill = _fill("5C1A1A")
    dup_fill  = _fill("3D2800")

    for ri, (idx, row_data) in enumerate(df.iterrows(), 2):
        ws.row_dimensions[ri].height = 18
        is_dup = idx in dup_rows
        for ci, (col, val) in enumerate(row_data.items(), 1):
            c = ws.cell(row=ri, column=ci)
            if pd.isna(val):
                c.value = "-"
                c.fill  = null_fill
                c.font  = _font(italic=True, color=C["text_dim"], size=10)
            else:
                c.value = val
                c.fill  = dup_fill if is_dup else _fill(C["bg_dark"] if ri % 2 == 0 else C["row_alt"])
                c.font  = _font(italic=True, color=C["gold"], size=10) if is_dup else _font(size=10)
            c.alignment = _left()
            c.border    = _border()

    ws.auto_filter.ref = ws.dimensions
    return ws


def build_kpi_sheet(wb, audit, df):
    ws = wb.create_sheet("KPI Dashboard")
    ws.sheet_properties.tabColor = "D4AC0D"
    for col_letter, width in zip("ABCDEFG", [28,18,18,18,18,18,18]):
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 40
    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "KEY PERFORMANCE INDICATORS"
    c.fill  = _fill(C["accent"])
    c.font  = Font(bold=True, size=18, color=C["white"], name="Calibri")
    c.alignment = _center()

    row = 3
    for ci, h in enumerate(["Column","Sum","Mean","Median","Std Dev","Min","Max"], 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = _fill(C["bg_header"])
        c.font = _font(bold=True, color=C["accent2"])
        c.alignment = _center()
        c.border = _border()
        ws.row_dimensions[row].height = 22

    row = 4
    for alt, (col, kpi) in enumerate(audit["kpis"].items()):
        ws.row_dimensions[row].height = 20
        bg = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
        for ci, val in enumerate([col, kpi["sum"], kpi["mean"], kpi["median"],
                                   kpi["std"], kpi["min"], kpi["max"]], 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.fill = _fill(bg)
            c.font = _font(size=10, bold=(ci == 1))
            c.alignment = _center() if ci > 1 else _left()
            c.border = _border()
            if ci > 1 and isinstance(val, (int, float)):
                c.number_format = '#,##0.00'
        row += 1

    if audit["salary_issues"]:
        row += 1
        write_section_title(ws, row, 1, "  SALARY ANALYSIS", 7)
        row += 1
        for ci, h in enumerate(["Column","Total","Mean","Median","Min","Max","Zeros"], 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.fill = _fill(C["gold"])
            c.font = _font(bold=True, color="1A0A0A")
            c.alignment = _center()
            c.border = _border()
            ws.row_dimensions[row].height = 22
        row += 1
        for col, info in audit["salary_issues"].items():
            ws.row_dimensions[row].height = 20
            for ci, val in enumerate([col, info["total"], info["mean"], info["median"],
                                       info["min"], info["max"], info["zero_count"]], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(C["bg_dark"])
                c.font = _font(size=10)
                c.alignment = _center() if ci > 1 else _left()
                c.border = _border()
                if ci > 1 and isinstance(val, (int, float)):
                    c.number_format = '#,##0.00'
            row += 1
    return ws


def build_error_sheet(wb, audit, df):
    ws = wb.create_sheet("Error Analysis")
    ws.sheet_properties.tabColor = "C0392B"
    for col_letter, width in zip("ABCDE", [22, 50, 16, 16, 30]):
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 38
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = "ERROR & ANOMALY ANALYSIS"
    c.fill  = _fill(C["accent"])
    c.font  = Font(bold=True, size=17, color=C["white"], name="Calibri")
    c.alignment = _center()

    row = 1

    def section(title):
        nonlocal row
        row += 1
        ws.row_dimensions[row].height = 26
        write_section_title(ws, row, 1, f"  {title}", 5)
        row += 1

    def hdr_row(headers):
        nonlocal row
        ws.row_dimensions[row].height = 20
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.fill = _fill(C["bg_header"])
            c.font = _font(bold=True, color=C["accent2"])
            c.alignment = _center()
            c.border = _border()
        row += 1

    def ok_row(text):
        nonlocal row
        c = ws.cell(row=row, column=1, value=text)
        c.fill = _fill(C["green"])
        c.font = _font(bold=True, color=C["white"])
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c.alignment = _left()
        row += 1

    # Missing
    section("MISSING VALUES")
    if audit["missing"]:
        hdr_row(["Column","Missing Count","% Missing","Severity","Recommendation"])
        for alt, (col, info) in enumerate(audit["missing"].items()):
            sev = "HIGH" if info["pct"] > 20 else "MEDIUM" if info["pct"] > 5 else "LOW"
            rec = "Drop column" if info["pct"] > 50 else "Impute median/mode" if info["pct"] > 10 else "Fill with default"
            bg  = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
            sev_c = C["red"] if sev=="HIGH" else C["orange"] if sev=="MEDIUM" else C["gold"]
            ws.row_dimensions[row].height = 18
            for ci, val in enumerate([col, info["count"], f"{info['pct']}%", sev, rec], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(sev_c) if ci == 4 else _fill(bg)
                c.font = _font(size=10, bold=(ci==4))
                c.alignment = _center() if ci != 5 else _left()
                c.border = _border()
            row += 1
    else:
        ok_row("No missing values detected")

    # Duplicates
    section("DUPLICATE ROWS")
    cnt = audit["duplicates"]["count"]
    c = ws.cell(row=row, column=1, value=f"Total Duplicates: {cnt}")
    c.fill = _fill(C["red"] if cnt > 0 else C["green"])
    c.font = _font(bold=True, color=C["white"], size=11)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c.alignment = _left()
    if audit["duplicates"]["rows"]:
        c2 = ws.cell(row=row, column=3,
                     value=f"Sample rows: {', '.join(map(str, audit['duplicates']['rows'][:5]))}")
        c2.fill = _fill(C["bg_dark"])
        c2.font = _font(size=10, italic=True)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        c2.alignment = _left()
    row += 1

    # Negative profits
    section("NEGATIVE PROFIT VALUES")
    if audit["negative_profits"]:
        for alt, (col, info) in enumerate(audit["negative_profits"].items()):
            bg = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
            ws.row_dimensions[row].height = 20
            for ci, val in enumerate([col, f"{info['count']} rows",
                                       f"Min: {info['min']:,.2f}", "HIGH",
                                       f"Rows: {', '.join(map(str, info['rows'][:5]))}"], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(C["red"]) if ci == 4 else _fill(bg)
                c.font = _font(size=10, bold=(ci==4))
                c.alignment = _left()
                c.border = _border()
            row += 1
    else:
        ok_row("No negative profit values detected")

    # Business logic
    section("BUSINESS LOGIC ERRORS")
    if audit["business_logic_errors"]:
        for alt, err in enumerate(audit["business_logic_errors"]):
            bg    = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
            sev_c = C["red"] if err["severity"]=="high" else C["orange"]
            ws.row_dimensions[row].height = 20
            for ci, val in enumerate([err["type"], err["description"],
                                       str(err["count"])+" rows",
                                       err["severity"].upper(),
                                       str(err["affected_rows"][:3])], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(sev_c) if ci == 4 else _fill(bg)
                c.font = _font(size=10, bold=(ci==4))
                c.alignment = _left()
                c.border = _border()
            row += 1
    else:
        ok_row("No business logic errors detected")

    # Anomalies
    section("STATISTICAL ANOMALIES  (Z-Score > 3)")
    if audit["statistical_anomalies"]:
        for alt, a in enumerate(audit["statistical_anomalies"]):
            bg = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
            vals_str = ", ".join(f"{v:,.2f}" for v in a["values"][:3])
            ws.row_dimensions[row].height = 20
            for ci, val in enumerate([a["column"], f"{a['count']} outliers",
                                       "Z > 3.0", "MEDIUM", vals_str], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(C["orange"]) if ci == 4 else _fill(bg)
                c.font = _font(size=10, bold=(ci==4))
                c.alignment = _left()
                c.border = _border()
            row += 1
    else:
        ok_row("No statistical anomalies detected")

    # Suspicious edits
    section("SUSPICIOUS MANUAL EDITS")
    if audit["suspicious_edits"]:
        for alt, s in enumerate(audit["suspicious_edits"]):
            bg = C["row_alt"] if alt % 2 == 0 else C["bg_dark"]
            ws.row_dimensions[row].height = 20
            for ci, val in enumerate([s["column"], s["description"],
                                       s["type"], "LOW", str(s["rows"][:5])], 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(C["gold"]) if ci == 4 else _fill(bg)
                c.font = _font(size=10, bold=(ci==4),
                               color="1A0A0A" if ci==4 else C["text_light"])
                c.alignment = _left()
                c.border = _border()
            row += 1
    else:
        ok_row("No suspicious edits detected")

    return ws


def build_chart_sheet(wb, df, audit):
    ws = wb.create_sheet("Charts")
    ws.sheet_properties.tabColor = "D4AC0D"
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "DATA VISUALIZATIONS"
    c.fill  = _fill(C["accent"])
    c.font  = Font(bold=True, size=18, color=C["white"], name="Calibri")
    c.alignment = _center()
    ws.row_dimensions[1].height = 38

    num_cols   = df.select_dtypes(include=[np.number]).columns.tolist()
    chart_row  = 3
    charts_done = 0

    if len(num_cols) >= 1 and len(df) >= 3:
        col_name = num_cols[0]
        ws.cell(row=chart_row, column=1, value="Index")
        ws.cell(row=chart_row, column=2, value=col_name)
        for i, val in enumerate(df[col_name].dropna().head(20).values, 1):
            ws.cell(row=chart_row+i, column=1, value=i)
            ws.cell(row=chart_row+i, column=2, value=float(val))
        chart = BarChart()
        chart.type    = "col"
        chart.title   = f"Distribution — {col_name}"
        chart.style   = 10
        chart.width   = 18
        chart.height  = 12
        data_ref = Reference(ws, min_col=2, min_row=chart_row,
                             max_row=chart_row+min(20, len(df)))
        chart.add_data(data_ref, titles_from_data=True)
        ws.add_chart(chart, f"D{chart_row}")
        chart_row  += 22
        charts_done += 1

    if audit["missing"]:
        pie_start = chart_row
        ws.cell(row=pie_start, column=1, value="Column")
        ws.cell(row=pie_start, column=2, value="Missing %")
        for i, (col, info) in enumerate(audit["missing"].items(), 1):
            ws.cell(row=pie_start+i, column=1, value=col)
            ws.cell(row=pie_start+i, column=2, value=info["pct"])
        pie = PieChart()
        pie.title  = "Missing Values Distribution"
        pie.style  = 10
        pie.width  = 16
        pie.height = 12
        data_ref   = Reference(ws, min_col=2, min_row=pie_start,
                               max_row=pie_start+len(audit["missing"]))
        labels_ref = Reference(ws, min_col=1, min_row=pie_start+1,
                               max_row=pie_start+len(audit["missing"]))
        pie.add_data(data_ref, titles_from_data=True)
        pie.set_categories(labels_ref)
        ws.add_chart(pie, f"D{chart_row}")
        chart_row  += 22
        charts_done += 1

    if len(num_cols) >= 2 and len(df) >= 5:
        col_name  = num_cols[1]
        line_start = chart_row
        ws.cell(row=line_start, column=1, value="Index")
        ws.cell(row=line_start, column=2, value=col_name)
        for i, val in enumerate(df[col_name].dropna().head(30).values, 1):
            ws.cell(row=line_start+i, column=1, value=i)
            ws.cell(row=line_start+i, column=2, value=float(val))
        line = LineChart()
        line.title   = f"Trend — {col_name}"
        line.style   = 10
        line.width   = 18
        line.height  = 12
        data_ref = Reference(ws, min_col=2, min_row=line_start,
                             max_row=line_start+min(30, len(df)))
        line.add_data(data_ref, titles_from_data=True)
        ws.add_chart(line, f"D{chart_row}")
        charts_done += 1

    if charts_done == 0:
        c = ws.cell(row=3, column=1, value="Not enough numeric data to generate charts.")
        c.font = _font(italic=True, color=C["text_dim"])
    return ws


def build_insights_sheet(wb, audit, df):
    ws = wb.create_sheet("Insights")
    ws.sheet_properties.tabColor = "27AE60"
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 65

    ws.row_dimensions[1].height = 38
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "INSIGHTS & RECOMMENDATIONS"
    c.fill  = _fill(C["accent"])
    c.font  = Font(bold=True, size=18, color=C["white"], name="Calibri")
    c.alignment = _center()

    score    = audit["health_score"]
    insights = []

    if score >= 85:
        insights.append(("OK", "Dataset Health", "Excellent — dataset is clean and ready for analysis."))
    elif score >= 65:
        insights.append(("WARN", "Dataset Health", f"Good — minor issues detected. Score: {score}%. Address warnings before reporting."))
    elif score >= 40:
        insights.append(("ERR", "Dataset Health", f"Poor — several issues need attention. Score: {score}%. Review all error flags."))
    else:
        insights.append(("ERR", "Dataset Health", f"Critical — serious integrity problems. Score: {score}%. Do not use without thorough cleansing."))

    if audit["duplicates"]["count"] > 0:
        insights.append(("WARN", "Duplicate Rows",
            f"Remove {audit['duplicates']['count']} duplicate rows using drop_duplicates() or Excel Remove Duplicates."))

    for col, info in audit["missing"].items():
        if info["pct"] > 50:
            insights.append(("ERR",  f"Drop '{col}'", f"{info['pct']}% missing. Consider dropping this column entirely."))
        elif info["pct"] > 10:
            insights.append(("WARN", f"Impute '{col}'", f"{info['pct']}% missing. Impute with median (numeric) or mode (categorical)."))
        else:
            insights.append(("INFO", f"Fill '{col}'", f"Only {info['pct']}% missing. Fill with a sensible default or forward-fill."))

    for col, info in audit["negative_profits"].items():
        insights.append(("ERR", f"Negative Profits: '{col}'",
            f"{info['count']} negative entries. Verify if returns/refunds or data entry errors."))

    for err in audit["business_logic_errors"]:
        insights.append(("ERR", "Logic Error",
            f"{err['description']}. Recalculate derived columns from source data."))

    if audit["statistical_anomalies"]:
        cols = [a["column"] for a in audit["statistical_anomalies"]]
        insights.append(("WARN", "Outliers Detected",
            f"Columns {', '.join(cols)} contain outliers (Z>3). Verify real data vs. input errors."))

    if audit["suspicious_edits"]:
        insights.append(("INFO", "Manual Edit Flags",
            f"{len(audit['suspicious_edits'])} columns flagged for unusual rounding patterns. Verify against source systems."))

    if not insights:
        insights.append(("OK", "All Clear", "Dataset passed all automated checks successfully."))

    icon_map  = {"OK": "OK",   "WARN": "WARN", "ERR": "ERR",  "INFO": "INFO"}
    color_map = {"OK": C["green"], "WARN": C["orange"], "ERR": C["red"], "INFO": C["gold"]}
    label_map = {"OK": "PASS", "WARN": "WARN", "ERR": "FAIL", "INFO": "NOTE"}

    row = 3
    for alt, (kind, title, text) in enumerate(insights):
        ws.row_dimensions[row].height = 36
        c1 = ws.cell(row=row, column=1, value=label_map[kind])
        c1.fill = _fill(color_map[kind])
        c1.font = Font(bold=True, size=9, color=C["white"], name="Calibri")
        c1.alignment = _center()
        c1.border = _border()

        c2 = ws.cell(row=row, column=2, value=title)
        c2.fill = _fill(C["row_alt"] if alt % 2 == 0 else C["bg_dark"])
        c2.font = _font(bold=True, size=11)
        c2.alignment = _left()
        c2.border = _border()

        c3 = ws.cell(row=row, column=3, value=text)
        c3.fill = _fill(C["row_alt"] if alt % 2 == 0 else C["bg_dark"])
        c3.font = _font(size=10)
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c3.border = _border()
        row += 1

    return ws


def generate_excel_report(df, audit):
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    build_cover_sheet(wb, audit, df)
    build_cleaned_data_sheet(wb, df, audit)
    build_kpi_sheet(wb, audit, df)
    build_error_sheet(wb, audit, df)
    build_chart_sheet(wb, df, audit)
    build_insights_sheet(wb, audit, df)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── routes ────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    with open("index.html", "rb") as f:
        return f.read(), 200, {"Content-Type": "text/html; charset=utf-8"}


@app.route("/analyze", methods=["POST"])
def analyze():
    try:
        if "file" not in request.files:
            return jsonify({"error": "No file uploaded"}), 400

        file     = request.files["file"]
        filename = file.filename or "upload"

        file_bytes = io.BytesIO(file.read())

        if filename.lower().endswith(".csv"):
            df = None
            for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
                try:
                    file_bytes.seek(0)
                    df = pd.read_csv(file_bytes, encoding=enc)
                    break
                except (UnicodeDecodeError, Exception):
                    continue
            if df is None:
                return jsonify({"error": "Could not decode CSV. Try saving as UTF-8."}), 400

        elif filename.lower().endswith(".xlsx"):
            try:
                file_bytes.seek(0)
                df = pd.read_excel(file_bytes, engine="openpyxl")
            except Exception as xe:
                return jsonify({"error": f"XLSX read error: {xe}"}), 400

        elif filename.lower().endswith(".xls"):
            try:
                file_bytes.seek(0)
                df = pd.read_excel(file_bytes, engine="xlrd")
            except Exception as xe:
                return jsonify({"error": f"XLS read error: {xe}"}), 400

        else:
            return jsonify({"error": "Only CSV, XLSX, or XLS files are supported."}), 400

        if df is None or df.empty:
            return jsonify({"error": "Uploaded file is empty or unreadable."}), 400

        audit = full_audit(df, filename)

        # ── store df + audit in memory (no disk) ──────────────────────────
        key = hashlib.md5(
            f"{filename}{datetime.now().isoformat()}{id(df)}".encode()
        ).hexdigest()[:16]

        # Convert df to JSON-serialisable records so it survives in _REPORT_STORE
        _REPORT_STORE[key] = {
            "records": df.to_dict(orient="list"),
            "columns": list(df.columns),
            "audit":   audit,
        }

        # Keep store lean – evict oldest if > 20 entries
        if len(_REPORT_STORE) > 20:
            oldest = next(iter(_REPORT_STORE))
            del _REPORT_STORE[oldest]

        summary = {
            "filename":             filename,
            "rows":                 audit["shape"]["rows"],
            "cols":                 audit["shape"]["cols"],
            "health_score":         audit["health_score"],
            "missing_count":        sum(v["count"] for v in audit["missing"].values()),
            "missing_cols":         len(audit["missing"]),
            "duplicate_count":      audit["duplicates"]["count"],
            "negative_profit_count":sum(v["count"] for v in audit["negative_profits"].values()),
            "logic_errors":         len(audit["business_logic_errors"]),
            "anomalies":            len(audit["statistical_anomalies"]),
            "suspicious":           len(audit["suspicious_edits"]),
            "columns":              audit["columns"][:20],
            "col_types":            dict(list(audit["col_types"].items())[:20]),
            "report_key":           key,
        }
        return jsonify(summary)

    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route("/download/<key>", methods=["GET"])
def download(key):
    try:
        if not re.match(r'^[a-f0-9]{16}$', key):
            return jsonify({"error": "Invalid key"}), 400

        entry = _REPORT_STORE.get(key)
        if entry is None:
            return jsonify({"error": "Report not found. Please re-upload your file."}), 404

        # Reconstruct DataFrame from stored records
        df    = pd.DataFrame(entry["records"], columns=entry["columns"])
        audit = entry["audit"]

        buf   = generate_excel_report(df, audit)
        fname = f"audit_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=fname,
        )

    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
